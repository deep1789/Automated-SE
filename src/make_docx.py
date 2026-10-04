"""Build a clean Word manuscript from the LaTeX sources (numbering taken from the compiled .aux, citations via citeproc)."""
import re, os, sys, subprocess, shutil
import pypandoc
P = "paper/"
os.makedirs(P + "docx_build", exist_ok=True)

# ---------------- numbering from main.aux
aux = open(P + "main.aux").read()
NUM = {m.group(1): m.group(2) for m in re.finditer(r"\\newlabel\{([^}]+)\}\{\{([^}]*)\}\{", aux)}
def ref(m): return NUM.get(m.group(1), "??")

def rd(f): return open(P + f).read()
def strip_comments(t): return re.sub(r"(?<!\\)%.*", "", t)

def inline_inputs(t):
    def rep(m):
        f = m.group(1); f = f if f.endswith(".tex") else f + ".tex"
        return inline_inputs(rd(f))
    return re.sub(r"\\input\{([^}]+)\}", rep, t)

# ---------------- body
secs = ["sec_intro", "sec_related", "sec_method", "sec_design", "sec_results", "sec_discussion", "sec_threats", "sec_conclusion"]
body = "\n\n".join(rd(s + ".tex") for s in secs)
body = strip_comments(body)

# TikZ figure -> PNG
body = re.sub(r"\\resizebox\{\\linewidth\}\{!\}\{\\input\{figs/fig_framework\.tex\}\}", r"\\includegraphics[width=15cm]{docx_build/fw.png}", body)
body = inline_inputs(body)

# tables: remove resizebox wrappers, add numbered captions
body = re.sub(r"\\resizebox\{\\linewidth\}\{!\}\{(\\begin\{tabular\}.*?\\end\{tabular\})\}", r"\1", body, flags=re.S)
def table_caption(m):
    cap, lab = m.group(1), m.group(2)
    return "\\caption{Table " + NUM.get(lab, "?") + ". " + cap + "}"
body = re.sub(r"\\caption\{((?:[^{}]|\{[^{}]*\}|\{(?:[^{}]|\{[^{}]*\})*\})*)\}\\label\{(tab:[^}]+)\}", table_caption, body)
def fig_caption(m):
    cap, lab = m.group(1), m.group(2)
    return "\\caption{Figure " + NUM.get(lab, "?") + ". " + cap + "}"
body = re.sub(r"\\caption\{((?:[^{}]|\{[^{}]*\}|\{(?:[^{}]|\{[^{}]*\})*\})*)\}\\label\{(fig:[^}]+)\}", fig_caption, body)
# include graphics: pdf -> png, widths
def incl(m):
    opts, path = m.group(1) or "", m.group(2)
    if path.endswith(".pdf"): path = path[:-4] + ".png"
    w = 15.0
    mm = re.search(r"width=([0-9.]*)\\linewidth", opts)
    if mm: w = 15.0 * (float(mm.group(1)) if mm.group(1) else 1.0)
    elif "width=15cm" in opts: w = 15.0
    return f"\\includegraphics[width={w:.1f}cm]{{{path}}}"
body = re.sub(r"\\includegraphics\[([^\]]*)\]\{([^}]+)\}", incl, body)
body = body.replace("\\centering", "")

# theorem-like environments -> numbered paragraphs (counters per type, document order)
cnt = {"lemma": 0, "proposition": 0, "corollary": 0}
NM = {"lemma": "Lemma", "proposition": "Proposition", "corollary": "Corollary"}
def thm(m):
    env, inner = m.group(1), m.group(2); cnt[env] += 1
    inner = re.sub(r"\\label\{[^}]*\}", "", inner)
    return f"\n\n\\textbf{{{NM[env]} {cnt[env]}.}} \\emph{{{inner.strip()}}}\n\n"
body = re.sub(r"\\begin\{(lemma|proposition|corollary)\}(.*?)\\end\{\1\}", thm, body, flags=re.S)
body = re.sub(r"\\begin\{proof\}(.*?)\\end\{proof\}", lambda m: "\n\n\\emph{Proof.} " + m.group(1).strip() + " $\\square$\n\n", body, flags=re.S)

# algorithm -> numbered plain list
ALG = r'''

\textbf{Algorithm 1.} Risk-controlled triage (calibration and deployment).

\emph{Input:} trained scorer $s(\cdot)$; labelled calibration set from the \emph{deployment} distribution; miss budget $\alpha$; optional PAC level $\delta$; optional group map $b(\cdot)$.

1. Collect the scores of the vulnerable calibration functions, sorted: $S^+$ ($n$ of them; per group if Mondrian).

2. Set $k=\lfloor\alpha(n+1)\rfloor$, or $k=j^\star(n,\alpha,\delta)$ for the PAC variant; if $k=0$, clear nothing.

3. Set the threshold $\hat\tau=s_{(k)}$ (clear iff $s(x)<\hat\tau$).

4. For each new function $x$: if $s(x)<\hat\tau$, auto-clear; otherwise send to review (or auto-flag if $s(x)\ge\hat\lambda$).

5. Monitor: on every batch of newly labelled outcomes, update the level with Eq. (ACI) or recalibrate on the most recent observed vulnerable functions.

'''
body = re.sub(r"\\begin\{algorithm\}.*?\\end\{algorithm\}", lambda m: ALG, body, flags=re.S)

# robust delimiters for Word
body = re.sub(r"\\#\\\{((?:[^{}]|\{[^{}]*\})*)\\\}", lambda m: "\\left|\\{" + m.group(1) + "\\}\\right|", body)
body = body.replace("\\lvert", "\\left|").replace("\\rvert", "\\right|")
# equations: labelled equations get explicit numbers; remove other labels
def eq_label(env_body):
    labs = re.findall(r"\\label\{(eq:[^}]+)\}", env_body)
    return labs
def eqenv(m):
    env, inner = m.group(1), m.group(2)
    lines = inner.split("\\\\")
    out = []
    for ln in lines:
        lab = re.search(r"\\label\{(eq:[^}]+)\}", ln)
        ln = re.sub(r"\\label\{[^}]*\}", "", ln)
        if lab: ln = ln.rstrip() + r"\qquad (" + NUM.get(lab.group(1), "?") + ")"
        out.append(ln)
    return "\\begin{" + env + "}" + "\\\\".join(out) + "\\end{" + env + "}"
body = re.sub(r"\\begin\{(equation|align)\}(.*?)\\end\{\1\}", eqenv, body, flags=re.S)
body = re.sub(r"\\eqref\{([^}]+)\}", lambda m: "(" + NUM.get(m.group(1), "??") + ")", body)
body = re.sub(r"\\(?:auto|c)?ref\{([^}]+)\}", ref, body)
body = re.sub(r"\\label\{[^}]*\}", "", body)

# run-in paragraph headings
body = re.sub(r"\\paragraph\{([^}]*)\}\s*", lambda m: "\n\n\\textbf{" + m.group(1).rstrip(".") + ".} ", body)
body = body.replace("\\begin{description}[leftmargin=2.2em,style=nextline,itemsep=2pt]", "\\begin{itemize}").replace("\\end{description}", "\\end{itemize}")
body = re.sub(r"\\item\[([^\]]*)\]", r"\\item \\textbf{\1} ", body)
body = re.sub(r"\\begin\{itemize\}\[[^\]]*\]", r"\\begin{itemize}", body)
body = body.replace("\\small", "").replace("\\footnotesize", "").replace("{\\scriptsize$\\pm$", "{$\\pm$")
body = re.sub(r"\\begin\{table\}\[t\]", r"\\begin{table}", body)


# ---------------- plain-text conversion of trivial inline math (numbers, symbols, Greek, single-letter variables)
GREEK = {"alpha": "α", "beta": "β", "gamma": "γ", "delta": "δ", "pi": "π", "tau": "τ", "rho": "ρ", "lambda": "λ", "epsilon": "ε", "sigma": "σ", "mu": "μ"}
SYM = {"%": "%", "times": "×", "pm": "±", "le": "≤", "leq": "≤", "ge": "≥", "geq": "≥", "to": "→", "approx": "≈", "lesssim": "≲", "sim": "~", "cdot": "·",
       "infty": "∞", "ne": "≠", "neq": "≠", ",": " ", ";": " ", "!": "", "ll": "≪", "gg": "≫", "rightarrow": "→", "square": "□"}
def simple_math(m):
    s = m.group(1)
    s = s.replace("{=}", "=").replace("\\%", "\\%")
    if re.search(r"[_^{}]|\\frac|\\sqrt|\\lfloor|\\Pr|\\mathrm|\\text|\\mathbb|\\bar|\\left|\\right|\\sum|\\max|\\in|\\mid", s): return m.group(0)
    def cmd(c):
        n = c.group(1)
        if n in GREEK: return GREEK[n]
        if n in SYM: return SYM[n]
        return "\\" + n + "\x00"      # unknown command -> abort
    out = re.sub(r"\\([a-zA-Z]+|[%,;!])", cmd, s)
    if "\\" in out or "\x00" in out: return m.group(0)
    out = out.replace("--", "–")
    out = re.sub(r"(?<![0-9.A-Za-z])-(?=[0-9])", "−", out)
    if re.search(r"[A-Za-z]{2,}", out): return m.group(0)
    if not re.fullmatch(r"[0-9A-Za-z.,:;%+\-−–=<>≤≥≈≲~×±→·∞≠≪≫□αβγδπτρλεσμ()/\s  |]*", out): return m.group(0)
    out = re.sub(r"(?<![A-Za-z\\])([A-Za-z])(?![A-Za-z])", lambda x: "\\emph{" + x.group(1) + "}", out)
    return out.replace("%", "\\%")
def simplify_math(txt):
    # protect display math
    parts = re.split(r"(\$\$.*?\$\$|\\begin\{(?:equation|align)\}.*?\\end\{(?:equation|align)\})", txt, flags=re.S)
    for i in range(0, len(parts), 2):
        parts[i] = re.sub(r"(?<!\\)\$([^$]+?)\$", simple_math, parts[i])
    return "".join(parts)

# ---------------- front matter
abs_src = strip_comments(rd("sec_abstract.tex"))
abs_txt = re.search(r"\\begin\{abstract\}(.*?)\\begin\{keyword\}", abs_src, flags=re.S).group(1).strip()
kw = re.search(r"\\begin\{keyword\}(.*?)\\end\{keyword\}", abs_src, flags=re.S).group(1).strip().replace("\\sep", ";")
abs_txt = re.sub(r"\\textbf\{([^}]*)\}", lambda m: "\n\n\\textbf{" + m.group(1) + "}", abs_txt)

front = r'''\begin{center}
\textbf{\Large Abstain Before You Trust: Risk-Controlled Triage for Automated Vulnerability Detection under Project, Temporal and Dataset Shift}

[Author names and affiliations to be added]

Manuscript prepared for the special issue ``Reliable and Trustworthy Automated Software Engineering'', Journal of Systems and Software
\end{center}

\section*{Abstract}

''' + abs_txt + "\n\n\\textbf{Keywords:} " + kw + "\n\n"

back = r'''
\section*{Data and code availability}
All code, split definitions, cached detector scores and the scripts that generate every table and figure are in the repository that accompanies this manuscript [repository URL and archived DOI to be added]. All three datasets are public (DiverseVul, Big-Vul, PrimeVul) and were obtained from the Hugging Face hub.

\section*{Declaration of generative AI and AI-assisted technologies in the writing process}
[To be confirmed and completed by the authors.] During the preparation of this work, an AI coding assistant ([tool name and version to be completed]) was used to help write the analysis code, run the experiments and draft the manuscript text. After using this tool, the author(s) reviewed and edited the content as needed and take(s) full responsibility for the content of the publication.

\section*{CRediT authorship contribution statement}
[To be added by the authors.]

\section*{Declaration of competing interest}
[To be added by the authors.]

\section*{Funding}
[To be added by the authors.]

\section*{References}

'''
body = simplify_math(body)
front = simplify_math(front)
doc = front + body + back
open(P + "docx_build/manuscript.tex", "w").write(doc)

# ---------------- pandoc
out = P + "docx_build/manuscript_raw.docx"
extra = ["--from=latex+raw_tex", "--to=docx", "--citeproc", f"--bibliography={P}refs.bib", f"--csl={P}docx_build/elsevier-harvard.csl",
         "--number-sections", f"--resource-path={P}", "--metadata=link-citations:true"]
pypandoc.convert_file(P + "docx_build/manuscript.tex", "docx", outputfile=out, extra_args=extra[1:] if False else ["-f", "latex", "--citeproc", f"--bibliography={P}refs.bib", f"--csl={P}docx_build/elsevier-harvard.csl", "--number-sections", f"--resource-path={P}:{P}docx_build"])
print("raw docx written", os.path.getsize(out))

# ---------------- post-processing (fonts, margins, table rules, captions, page numbers)
from docx import Document
from docx.shared import Pt, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
d = Document(out)
for s in d.sections:
    s.left_margin = s.right_margin = Cm(2.3); s.top_margin = s.bottom_margin = Cm(2.3)
    s.page_width, s.page_height = Cm(21.0), Cm(29.7)
def setfont(style, size=None, bold=None, italic=None):
    f = style.font; f.name = "Times New Roman"
    rpr = style.element.get_or_add_rPr(); rf = rpr.find(qn("w:rFonts"))
    if rf is None: rf = OxmlElement("w:rFonts"); rpr.append(rf)
    for a in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"): rf.set(qn(a), "Times New Roman")
    for a in ("w:asciiTheme", "w:hAnsiTheme", "w:cstheme", "w:eastAsiaTheme"):
        if rf.get(qn(a)) is not None: del rf.attrib[qn(a)]
    if size: f.size = Pt(size)
    if bold is not None: f.bold = bold
    if italic is not None: f.italic = italic
    f.color.rgb = None if False else f.color.rgb
for st in d.styles:
    try:
        n = st.name
        if n in ("Normal", "Body Text", "First Paragraph", "Compact", "Block Text", "Abstract", "Author", "Date"):
            setfont(st, 11); 
            if n in ("Body Text", "First Paragraph", "Normal"):
                st.paragraph_format.space_after = Pt(6); st.paragraph_format.line_spacing = 1.15; st.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        elif n.startswith("Heading"):
            lv = int(n.split()[-1]) if n.split()[-1].isdigit() else 1
            setfont(st, {1: 14, 2: 12.5, 3: 11.5}.get(lv, 11), bold=True, italic=(lv >= 3)); st.font.color.rgb = None
            from docx.shared import RGBColor; st.font.color.rgb = RGBColor(0, 0, 0)
            st.paragraph_format.space_before = Pt(12 if lv == 1 else 8); st.paragraph_format.space_after = Pt(4); st.paragraph_format.keep_with_next = True
        elif n in ("Table Caption", "Image Caption", "Captioned Figure"):
            setfont(st, 9.5); st.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY if n != "Captioned Figure" else WD_ALIGN_PARAGRAPH.CENTER
            if n == "Table Caption": st.paragraph_format.keep_with_next = True
    except Exception as e:
        pass
# make caption prefix ("Table N." / "Figure N.") bold
for p in d.paragraphs:
    if p.style.name in ("Table Caption", "Image Caption"):
        m = re.match(r"^((?:Table|Figure) \d+\.)", p.text)
        if m and p.runs:
            r0 = p.runs[0]
            if r0.text.startswith(m.group(1)):
                rest = r0.text[len(m.group(1)):]; r0.text = m.group(1); r0.bold = True
                new = p.add_run(rest); new.bold = False
                # move the remainder run right after the bold run
                r0._r.addnext(new._r)
    if p.style.name == "Captioned Figure" or p.runs and p._p.xpath(".//pic:pic"):
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
# tables: three-rule (booktabs-like) look, small font
def border(el, tag, sz=8, val="single"):
    b = OxmlElement(tag); b.set(qn("w:val"), val); b.set(qn("w:sz"), str(sz)); b.set(qn("w:space"), "0"); b.set(qn("w:color"), "000000"); el.append(b)
for t in d.tables:
    ncol = len(t.columns); fs = 9 if ncol <= 5 else (8 if ncol <= 8 else (7 if ncol <= 11 else 6))
    tblPr = t._tbl.tblPr
    for old in tblPr.findall(qn("w:tblBorders")): tblPr.remove(old)
    tb = OxmlElement("w:tblBorders"); border(tb, "w:top", 12); border(tb, "w:bottom", 12); tblPr.append(tb)
    tw = tblPr.find(qn("w:tblW"))
    if tw is None: tw = OxmlElement("w:tblW"); tblPr.append(tw)
    tw.set(qn("w:type"), "pct"); tw.set(qn("w:w"), "5000")
    lay = tblPr.find(qn("w:tblLayout"))
    if lay is None: lay = OxmlElement("w:tblLayout"); tblPr.append(lay)
    lay.set(qn("w:type"), "autofit")
    for ri, row in enumerate(t.rows):
        for c in row.cells:
            tcPr = c._tc.get_or_add_tcPr()
            if ri == 0:
                bd = OxmlElement("w:tcBorders"); border(bd, "w:bottom", 6); tcPr.append(bd)
            mar = OxmlElement("w:tcMar")
            for side in ("left", "right"):
                e = OxmlElement(f"w:{side}"); e.set(qn("w:w"), "40"); e.set(qn("w:type"), "dxa"); mar.append(e)
            tcPr.append(mar)
            for p in c.paragraphs:
                p.paragraph_format.space_after = Pt(0); p.paragraph_format.space_before = Pt(0); p.paragraph_format.line_spacing = 1.0
                p.alignment = WD_ALIGN_PARAGRAPH.LEFT if c is row.cells[0] else WD_ALIGN_PARAGRAPH.CENTER
                for r in p.runs: r.font.size = Pt(fs); r.font.name = "Times New Roman"
                if ri == 0:
                    for r in p.runs: r.bold = True
# footer page numbers
for s in d.sections:
    f = s.footer.paragraphs[0] if s.footer.paragraphs else s.footer.add_paragraph(); f.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for tag, txt in (("begin", None), (None, "PAGE"), ("end", None)):
        r = f.add_run()
        if tag: fc = OxmlElement("w:fldChar"); fc.set(qn("w:fldCharType"), tag); r._r.append(fc)
        else: it = OxmlElement("w:instrText"); it.set(qn("xml:space"), "preserve"); it.text = txt; r._r.append(it)
# enforce OOXML child order inside tblPr (Word is strict)
ORDER = ["tblStyle", "tblpPr", "tblOverlap", "bidiVisual", "tblStyleRowBandSize", "tblStyleColBandSize", "tblW", "jc", "tblCellSpacing", "tblInd", "tblBorders", "shd", "tblLayout", "tblCellMar", "tblLook", "tblCaption", "tblDescription"]
for t_ in d.tables:
    pr = t_._tbl.tblPr; kids = list(pr)
    kids.sort(key=lambda e: ORDER.index(e.tag.split("}")[1]) if e.tag.split("}")[1] in ORDER else 99)
    for k in list(pr): pr.remove(k)
    for k in kids: pr.append(k)
TC_ORDER = ["cnfStyle", "tcW", "gridSpan", "hMerge", "vMerge", "tcBorders", "shd", "noWrap", "tcMar", "textDirection", "tcFitText", "vAlign", "hideMark"]
for t_ in d.tables:
    for row in t_.rows:
        for c in row.cells:
            pr = c._tc.tcPr
            if pr is None: continue
            kids = list(pr); kids.sort(key=lambda e: TC_ORDER.index(e.tag.split("}")[1]) if e.tag.split("}")[1] in TC_ORDER else 99)
            for k in list(pr): pr.remove(k)
            for k in kids: pr.append(k)
final = "paper/Manuscript_clean.docx"
d.save(final); print("saved", final, os.path.getsize(final))
